import sqlite3
import openpyxl
import os

EXCEL_FILE = "fe_study_master_23sheets.xlsx"
DB_FILE = "fe_study.db"
IMAGE_DIR = "static/images"

if not os.path.exists(EXCEL_FILE):
    print("❌ Excelファイルが見つかりません。")
    exit(1)

# 画像マップ
image_map = {}
if os.path.exists(IMAGE_DIR):
    for f in os.listdir(IMAGE_DIR):
        base, ext = os.path.splitext(f)
        if ext.lower() in [".png", ".jpg", ".jpeg"]:
            image_map[base] = os.path.join(IMAGE_DIR, f)

conn = sqlite3.connect(DB_FILE)
cur = conn.cursor()

def safe_drop(target_name):
    cur.execute("SELECT type FROM sqlite_master WHERE name = ?", (target_name,))
    row = cur.fetchone()
    if row:
        target_type = row[0].upper()
        cur.execute(f"DROP {target_type} {target_name};")

safe_drop("terms")
safe_drop("terms_master")

# 1. 用語学習用テーブル (terms) の新規作成
cur.execute("""
CREATE TABLE terms (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    term_name TEXT UNIQUE NOT NULL,
    category TEXT,
    sub_category TEXT,
    metaphor_def TEXT,         -- 💡 直感理解（日常の比喩）
    practical_ex TEXT,         -- 🌍 実社会での活用例
    official_def TEXT,         -- 公式定義
    exam_point_text TEXT,      -- 📖 試験対策
    url TEXT,
    image_path TEXT
);
""")

# 2. 問題演習用テーブル (terms_master) の新規作成
cur.execute("""
CREATE TABLE terms_master (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    term_name TEXT,
    category TEXT,
    sub_category TEXT,
    question_text TEXT,
    choice_a TEXT,
    choice_b TEXT,
    choice_c TEXT,
    choice_d TEXT,
    correct_ans TEXT,
    explanation_beginner TEXT,
    explanation_advanced TEXT
);
""")

wb = openpyxl.load_workbook(EXCEL_FILE, data_only=True)
term_count = 0
quiz_count = 0

for sheetname in wb.sheetnames:
    if sheetname == "全用語一覧":
        continue

    ws = wb[sheetname]
    headers = [cell.value for cell in ws[1] if cell.value is not None]
    
    if "term_name" not in headers:
        continue

    def get_val(r_idx, col_name):
        if col_name in headers:
            c_idx = headers.index(col_name) + 1
            v = ws.cell(row=r_idx, column=c_idx).value
            return str(v).strip() if v is not None and str(v).strip() not in ["None", "nan"] else ""
        return ""

    for r in range(2, ws.max_row + 1):
        term = get_val(r, "term_name")
        if not term:
            continue

        cat = get_val(r, "category")
        subcat = get_val(r, "sub_category") or sheetname
        meta = get_val(r, "intuitive_def_final") or get_val(r, "intuitive_def")
        prac = get_val(r, "practical_example_final") or get_val(r, "practical_example")
        off = get_val(r, "kakomon_official_def") or get_val(r, "official_def")
        exam = get_val(r, "exam_point_text")
        url = get_val(r, "kakomon_url") or get_val(r, "url")
        img = image_map.get(term, "")

        # terms へ挿入（最新の比喩・実用例で上書き）
        cur.execute("""
        INSERT OR REPLACE INTO terms (
            term_name, category, sub_category, metaphor_def, practical_ex,
            official_def, exam_point_text, url, image_path
        ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)
        """, (term, cat, subcat, meta, prac, off, exam, url, img))
        term_count += 1

        # 問題データがあれば terms_master へ挿入
        q_text = get_val(r, "question_text")
        if q_text:
            cur.execute("""
            INSERT INTO terms_master (
                term_name, category, sub_category, question_text,
                choice_a, choice_b, choice_c, choice_d, correct_ans,
                explanation_beginner, explanation_advanced
            ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """, (
                term, cat, subcat, q_text,
                get_val(r, "choice_a"), get_val(r, "choice_b"),
                get_val(r, "choice_c"), get_val(r, "choice_d"),
                get_val(r, "correct_ans"),
                get_val(r, "exp_beginner_final") or get_val(r, "explanation_beginner"),
                get_val(r, "exp_advanced_final") or get_val(r, "explanation_advanced")
            ))
            quiz_count += 1

conn.commit()
conn.close()
print(f"🎉 データベース再構築完了！")
print(f"・用語データ (terms): {term_count} 件")
print(f"・問題演習データ (terms_master): {quiz_count} 件")
