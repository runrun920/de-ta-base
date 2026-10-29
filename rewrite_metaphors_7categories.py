import os
import sys
import time
import signal
from pydantic import BaseModel, Field
import openpyxl
from google import genai
from google.genai import types

CREDENTIALS_JSON = "gcp_key.json"
PROJECT_ID = "gen-lang-client-0984224626"
LOCATION = "us-central1"
EXCEL_FILE = "fe_study_master_23sheets.xlsx"

if os.path.exists(CREDENTIALS_JSON):
    os.environ["GOOGLE_APPLICATION_CREDENTIALS"] = os.path.abspath(CREDENTIALS_JSON)

# 厳選した重要7分野
TARGET_SHEETS = [
    "3. コンピュータ構成要素",
    "5. ソフトウェア",
    "9. データベース",
    "10. ネットワーク",
    "11. セキュリティ",
    "12. システム開発技術",
    "22. 企業活動"
]

class SuperIntuitiveContent(BaseModel):
    intuitive_def: str = Field(
        description="専門用語（テーブル、データ、パケット、暗号など）を一切使わず、料理・学校・ゲーム・コンビニなどの日常のモノや行動だけで例えた直感的な比喩説明（2〜3文）"
    )
    practical_example: str = Field(
        description="スマホ、銀行ATM、LINE、Amazon、Suicaなど、誰もが知る身近なサービスでどう使われているか具体的にイメージできる解説（2〜3文）"
    )

try:
    client = genai.Client(vertexai=True, project=PROJECT_ID, location=LOCATION)
    print("🔑 Google Cloud Vertex AI に接続しました。")
except Exception as e:
    print(f"❌ 認証エラー: {e}")
    sys.exit(1)

wb_global = None

def save_wb():
    global wb_global
    if wb_global:
        wb_global.save(EXCEL_FILE)

def signal_handler(sig, frame):
    print("\n\n🛑 中断シグナルを受信しました。直前の用語まで安全に保存して終了します...")
    save_wb()
    sys.exit(0)

signal.signal(signal.SIGINT, signal_handler)

def generate_better_metaphor(term_name: str, sub_cat: str, off_def: str):
    prompt = f"""
あなたは教育心理学とIT初等教育の最高のエキスパートです。
以下のIT用語について、ITの知識がまったくない文系大学生が「あ、そういうことか！」と一瞬で膝を打つような、
【最高品質の日常の比喩】と【身近な実社会の活用例】を作成してください。

【対象用語】: {term_name}
【分野】: {sub_cat}
【公式定義】: {off_def}

---
■ 厳格なルール：
1. 【直感理解（日常の比喩）】：
   - コンピュータ、データ、システム、テーブル、暗号、パケットなどのIT用語・専門用語を【完全に禁止】します。
   - レストランの厨房、学校の出欠、図書館の本、スーパーのレジ、ゲームのセーブ機能など、誰もが知っている日常のシーンに例えてください。
   - 「〜のようなものです。たとえば〜」の形式でわかりやすく書いてください。

2. 【実社会での活用例】：
   - 仕様書のような難解な説明（〜の整合性を保つ、〜のプロトコルに準拠するなど）は【完全に禁止】します。
   - 「メルカリで出品するとき」「LINEでメッセージを送るとき」「銀行ATMでお金を引き出すとき」「Suicaで改札を通るとき」など、具体的なサービス名を出して、初学者が利用者の視点でイメージできるように書いてください。
"""
    try:
        res = client.models.generate_content(
            model="gemini-2.5-flash",
            contents=prompt,
            config=types.GenerateContentConfig(
                response_mime_type="application/json",
                response_schema=SuperIntuitiveContent,
                temperature=0.3
            )
        )
        return res.parsed
    except Exception as e:
        print(f"\n❌ APIエラー ({term_name}): {e}")
        return None

def main():
    global wb_global

    if not os.path.exists(EXCEL_FILE):
        print(f"❌ {EXCEL_FILE} が見つかりません。")
        return

    print(f"📖 {EXCEL_FILE} を読み込み中...")
    wb_global = openpyxl.load_workbook(EXCEL_FILE)

    total_updated = 0

    for sheetname in TARGET_SHEETS:
        if sheetname not in wb_global.sheetnames:
            continue

        ws = wb_global[sheetname]
        headers = [cell.value for cell in ws[1] if cell.value is not None]

        t_col = headers.index("term_name") + 1
        meta_col = (headers.index("intuitive_def_final") + 1) if "intuitive_def_final" in headers else None
        prac_col = (headers.index("practical_example_final") + 1) if "practical_example_final" in headers else None
        off_col = (headers.index("kakomon_official_def") + 1) if "kakomon_official_def" in headers else None

        if not meta_col or not prac_col:
            print(f"⚠️ シート '{sheetname}' に必要なカラムがありません。スキップします。")
            continue

        print(f"\n==========================================")
        print(f"🚀 比喩・実用例の超高品質リライト開始: 【{sheetname}】")
        print(f"==========================================")

        for r in range(2, ws.max_row + 1):
            term = ws.cell(row=r, column=t_col).value
            if not term or str(term).strip() in ["", "nan"]:
                continue

            term = str(term).strip()
            off_def = str(ws.cell(row=r, column=off_col).value or "").strip() if off_col else ""

            print(f"[{r-1}/{ws.max_row-1}] リライト中: {term} ...", end="", flush=True)

            data = generate_better_metaphor(term, sheetname, off_def)
            if data:
                ws.cell(row=r, column=meta_col, value=data.intuitive_def)
                ws.cell(row=r, column=prac_col, value=data.practical_example)
                total_updated += 1
                save_wb() # 1件ごとに即時セーブ
                print(" ✅ 完了・保存済")
            else:
                print(" ⚠️ スキップ")

            time.sleep(0.6)

    print(f"\n🎉 7分野の比喩・実用例のリライトがすべて完了しました！（更新件数: {total_updated} 件）")

if __name__ == "__main__":
    main()
