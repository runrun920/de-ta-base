import os
import sys
import time
from typing import Literal
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
import pandas as pd
from pydantic import BaseModel, Field
from google import genai
from google.genai import types

CREDENTIALS_JSON = "gcp_key.json"
PROJECT_ID = "gen-lang-client-0984224626"
LOCATION = "us-central1"
MASTER_FILE = "fe_study_master_23sheets.xlsx"
EVAL_FILE = "比喩・実用例_品質精査評価シート.xlsx"

if os.path.exists(CREDENTIALS_JSON):
    os.environ["GOOGLE_APPLICATION_CREDENTIALS"] = os.path.abspath(CREDENTIALS_JSON)

# リライト対象の全7分野をチェック対象に設定
TARGET_SHEETS = [
    "9. データベース",
    "11. セキュリティ",
    "3. コンピュータ構成要素",
    "5. ソフトウェア",
    "10. ネットワーク",
    "12. システム開発技術",
    "22. 企業活動"
]

HEADERS = [
    "No", "用語名", "中分類",
    "直感理解（日常の比喩）", "実社会での活用例",
    "【評価1】比喩の身近さ", 
    "【評価2】実用例の具体性",
    "【評価3】技術正確性",
    "指摘事項・改善案コメント"
]

class StrictMetaphorEvaluation(BaseModel):
    eval1_metaphor: Literal["○", "△", "×"] = Field(description="直感理解の身近さ。○(日常の例えとして秀逸), △(やや難解・抽象的), ×(比喩になっていない)")
    eval2_example: Literal["○", "△", "×"] = Field(description="実用例のわかりやすさ。○(具体的でイメージしやすい), △(仕様書的・難解), ×(不適切)")
    eval3_accuracy: Literal["○", "△", "×"] = Field(description="概念の正確性。○(誤解なし), △(軽微なズレあり), ×(致命的な誤謬あり)")
    comment: str = Field(description="初学者に伝わるようにするための具体的改善提案（60〜100文字程度）")

try:
    client = genai.Client(vertexai=True, project=PROJECT_ID, location=LOCATION)
except Exception as e:
    print(f"❌ Vertex AI 接続エラー: {e}")
    sys.exit(1)

# 評価シートの読み込み（既存のデータベースやセキュリティの完了データを引き継ぐ）
if os.path.exists(EVAL_FILE):
    wb_eval = openpyxl.load_workbook(EVAL_FILE)
else:
    wb_eval = openpyxl.Workbook()
    wb_eval.remove(wb_eval.active)

header_fill = PatternFill(start_color="3B5998", end_color="3B5998", fill_type="solid")
header_font = Font(name="Yu Gothic", size=10, bold=True, color="FFFFFF")
body_font = Font(name="Yu Gothic", size=9)
thin_side = Side(border_style="thin", color="CCCCCC")
cell_border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)
align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)
align_left = Alignment(horizontal="left", vertical="top", wrap_text=True)

excel_master = pd.ExcelFile(MASTER_FILE)

try:
    for sheetname in TARGET_SHEETS:
        if sheetname not in excel_master.sheet_names:
            continue

        eval_sheet_title = f"{sheetname} 査読"
        print(f"\n==========================================")
        print(f"🧐 査読セッション: {eval_sheet_title}")
        print(f"==========================================")

        df = pd.read_excel(MASTER_FILE, sheet_name=sheetname)
        meta_col = "intuitive_def_final" if "intuitive_def_final" in df.columns else "intuitive_def"
        prac_col = "practical_example_final" if "practical_example_final" in df.columns else "practical_example"
        
        valid_df = df[df[meta_col].notna() | df[prac_col].notna()].copy()
        all_rows = valid_df.to_dict("records")
        total_items = len(all_rows)

        # シートの存在確認・初期化（未作成分野は新規シート作成）
        if eval_sheet_title in wb_eval.sheetnames:
            ws = wb_eval[eval_sheet_title]
        else:
            ws = wb_eval.create_sheet(title=eval_sheet_title)
            ws.row_dimensions[1].height = 28
            for col_idx, h_text in enumerate(HEADERS, start=1):
                c = ws.cell(row=1, column=col_idx, value=h_text)
                c.fill = header_fill
                c.font = header_font
                c.alignment = align_center
                c.border = cell_border

        # 既存シートから「既に精査済みの用語名」を取得（二重査読を防止）
        existing_terms = set()
        for r in range(2, ws.max_row + 1):
            t_val = ws.cell(row=r, column=2).value # 2列目: 用語名
            e1_val = ws.cell(row=r, column=6).value # 6列目: 評価1
            if t_val and e1_val:
                existing_terms.add(str(t_val).strip())

        print(f"📊 登録済み用語数: {len(existing_terms)} / {total_items} 件（残り: {total_items - len(existing_terms)} 件）")

        for i, r_data in enumerate(all_rows):
            term = str(r_data.get("term_name", "")).strip()
            if not term:
                continue

            # 既に精査済みの用語はスキップ
            if term in existing_terms:
                continue

            subcat = str(r_data.get("sub_category", sheetname)).strip()
            meta = str(r_data.get(meta_col, "")).strip()
            prac = str(r_data.get(prac_col, "")).strip()

            print(f"[{i+1}/{total_items}] 査読中: {term}...", end="", flush=True)

            prompt = f"""
あなたは教育心理学および情報科学教育の専門家です。
初学者がITの抽象概念を直感的に理解できるよう作成された「日常の比喩（直感理解）」と「実社会での活用例」を厳格に査読してください。

【対象用語】: {term}
【分野】: {subcat}
【直感理解（日常の比喩）】: {meta}
【実社会での活用例】: {prac}

### 査読手順：
1. 直感理解が、料理や日常の行動などの「親しみやすい例え」になっているか検証する。専門用語を使ってそのまま説明している場合は×または△。
2. 実社会の例が、スマホ、銀行ATM、ECサイトなど「初学者がイメージできる身近な使われ方」になっているか検証する。
3. 判定記号は必ず「○」「△」「×」の1文字のみを出力すること。
"""

            time.sleep(1.0)
            try:
                res = client.models.generate_content(
                    model="gemini-2.5-flash",
                    contents=prompt,
                    config=types.GenerateContentConfig(
                        response_mime_type="application/json",
                        response_schema=StrictMetaphorEvaluation,
                        temperature=0.1
                    )
                )
                ev: StrictMetaphorEvaluation = res.parsed
                e1, e2, e3, comm = ev.eval1_metaphor, ev.eval2_example, ev.eval3_accuracy, ev.comment
            except Exception as ex:
                e1, e2, e3, comm = "エラー", "エラー", "エラー", f"APIエラー: {str(ex)[:60]}"

            clean_e1 = e1[0] if e1 in ["○", "△", "×"] else ("○" if "○" in e1 else ("△" if "△" in e1 else "×"))
            clean_e2 = e2[0] if e2 in ["○", "△", "×"] else ("○" if "○" in e2 else ("△" if "△" in e2 else "×"))
            clean_e3 = e3[0] if e3 in ["○", "△", "×"] else ("○" if "○" in e3 else ("△" if "△" in e3 else "×"))

            # 次の空き行（末尾）に追記
            next_row = ws.max_row + 1
            ws.row_dimensions[next_row].height = 65

            vals = [next_row - 1, term, subcat, meta, prac, clean_e1, clean_e2, clean_e3, comm]
            for col_idx, val in enumerate(vals, start=1):
                cell = ws.cell(row=next_row, column=col_idx, value=val)
                cell.font = body_font
                cell.border = cell_border
                if col_idx in [1, 6, 7, 8]:
                    cell.alignment = align_center
                else:
                    cell.alignment = align_left

            # 1件ごとに即座にExcelへ保存
            wb_eval.save(EVAL_FILE)
            existing_terms.add(term)
            print(f" ✅ 完了・保存済 ({clean_e1}/{clean_e2}/{clean_e3})")

        # 列幅の整形
        col_widths = {1: 6, 2: 18, 3: 16, 4: 38, 5: 38, 6: 12, 7: 12, 8: 12, 9: 45}
        for col_idx, width in col_widths.items():
            col_letter = openpyxl.utils.get_column_letter(col_idx)
            ws.column_dimensions[col_letter].width = width

        wb_eval.save(EVAL_FILE)

except KeyboardInterrupt:
    print("\n\n🛑 ユーザーによる中断を検知しました。")
    print(f"💾 直前に完了した用語まで '{EVAL_FILE}' に確実に保存されています。")
    print("次回実行時は、未精査の用語から自動的に再開されます。")
    sys.exit(0)

print(f"\n🎉 7分野すべての用語の品質精査が完了しました！ -> {EVAL_FILE}")
