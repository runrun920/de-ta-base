import os
import sqlite3
import openpyxl

EXCEL_FILE = "fe_study_master_23sheets.xlsx"
DB_FILE = "fe_study.db"
IMAGE_DIR = "static/images"

if not os.path.exists(EXCEL_FILE) or not os.path.exists(DB_FILE):
    print("❌ 必要なファイルが見つかりません。")
    exit(1)

print("⏳ Excelとローカル画像をスキャンしてDBに同期中...")
wb = openpyxl.load_workbook(EXCEL_FILE, data_only=True)
conn = sqlite3.connect(DB_FILE)
cur = conn.cursor()

# 存在する画像ファイルのマップを作成（.png, .jpg, .jpeg 対応）
existing_images = {}
if os.path.exists(IMAGE_DIR):
    for f in os.listdir(IMAGE_DIR):
        base, ext = os.path.splitext(f)
        if ext.lower() in [".png", ".jpg", ".jpeg"]:
            existing_images[base] = os.path.join(IMAGE_DIR, f)

updated_count = 0
for sheetname in wb.sheetnames:
    ws = wb[sheetname]
    headers = [cell.value for cell in ws[1]]
    
    if "term_name" not in headers:
        continue
    
    term_idx = headers.index("term_name") + 1
    exam_idx = headers.index("exam_point_text") + 1 if "exam_point_text" in headers else None
    
    for r in range(2, ws.max_row + 1):
        term = ws.cell(row=r, column=term_idx).value
        text = ws.cell(row=r, column=exam_idx).value if exam_idx else None
        
        if not term:
            continue
        
        term = str(term).strip()
        img_path = existing_images.get(term, None)
        
        # テキストまたは画像が存在する場合にDBを更新
        if text or img_path:
            updates = []
            params = []
            
            if text:
                updates.append("intuitive_def = ?")
                params.append(str(text).strip())
            
            if img_path:
                updates.append("image_path = ?")
                params.append(img_path)
            
            params.append(term)
            query = f"UPDATE terms SET {', '.join(updates)} WHERE term_name = ?;"
            cur.execute(query, params)
            if cur.rowcount > 0:
                updated_count += 1

conn.commit()
conn.close()

print(f"🎉 同期完了！ 合計 {updated_count} 件の用語データをDB（fe_study.db）に反映しました。")
