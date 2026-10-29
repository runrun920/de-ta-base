import openpyxl
import sqlite3
import os

EXCEL_FILE = "fe_study_master_23sheets.xlsx"
DB_FILE = "fe_study.db"

# ここに用語ごとの試験対策テキストを追加・蓄積していきます
TEXT_DATA = {
    "DDL": """【試験の決め手（正解を見抜くキーワード）】
・スキーマ、テーブル、ビュー、インデックスなどの「枠組み・構造」の定義・作成・変更・削除
・CREATE、DROP、ALTER

【誤答肢に惑わされないポイント】
データの検索（SELECT）や更新（INSERT、UPDATE、DELETE）はDML（データ操作言語）の役割です。「レコード（行）の追加や削除をする」という選択肢は代表的な誤答（不正解）であり、DDLはあくまでデータを入れるための「器（テーブル構造そのもの）」を定義・操作するものであることを見抜いてください。

【概念理解・仕組み（図解連動）】
データを保存する「データベースの枠組み（スキーマやテーブル）」を設計し、システムにその形を定義・作成・変更・削除するための専用コマンド群です。
・なぜ必要か：データ（中身）を入れる前に、まず「どんな名前のテーブルにし、どんな種類の列（データ型）を持たせるか」という構造（枠組み）を定義する手段が必要だからです。
・仕組み：設計者が「CREATE（創る）」「DROP（壊す・削除する）」「ALTER（変更する）」といった専用のDDL命令をDBMSに送り、データベースの物理的・論理的な構造を構築・編集します。
・最大の強み：データの中身（レコード）を一切触らずに、「テーブルそのものを新規作成する」「使わなくなったテーブルを消去する」「新しい列（カラム）をテーブルに追加する」といった全体設計の管理を行えます。"""
}

# 1. Excelファイルの更新
if os.path.exists(EXCEL_FILE):
    print("⏳ Excelファイルを更新中...")
    wb = openpyxl.load_workbook(EXCEL_FILE)
    excel_updated = 0
    for sheetname in wb.sheetnames:
        ws = wb[sheetname]
        headers = [cell.value for cell in ws[1]]
        if "term_name" not in headers or "exam_point_text" not in headers:
            continue
        
        t_col = headers.index("term_name") + 1
        e_col = headers.index("exam_point_text") + 1
        
        for r in range(2, ws.max_row + 1):
            term = str(ws.cell(row=r, column=t_col).value or "").strip()
            if term in TEXT_DATA:
                ws.cell(row=r, column=e_col, value=TEXT_DATA[term])
                excel_updated += 1
                
    wb.save(EXCEL_FILE)
    print(f"✅ Excel更新完了: 延べ {excel_updated} 箇所にテキストを書き込みました。")

# 2. SQLiteデータベース（fe_study.db）の更新
if os.path.exists(DB_FILE):
    print("⏳ DBを更新中...")
    conn = sqlite3.connect(DB_FILE)
    cur = conn.cursor()
    db_updated = 0
    for term, text in TEXT_DATA.items():
        cur.execute("UPDATE terms SET intuitive_def = ? WHERE term_name = ?;", (text, term))
        if cur.rowcount > 0:
            db_updated += cur.rowcount
    conn.commit()
    conn.close()
    print(f"🎉 DB更新完了: {db_updated} 件の用語の解説テキストを反映しました。")

